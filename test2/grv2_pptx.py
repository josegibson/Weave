import re
import openpyxl
from pptx import Presentation
from openpyxl.utils import column_index_from_string
import tempfile
import os
from pptx.util import Inches

def log_warning(message):
    print(f"WARNING: {message}")

def log_info(message):
    print(f"INFO: {message}")

def find_sheet(workbook, sheet_name):
    """Find a sheet in the workbook, normalizing spaces and handling case-insensitive matches."""
    # Normalize the input sheet name: strip spaces, collapse multiple spaces
    normalized_input = re.sub(r'\s+', ' ', sheet_name.strip()).lower()
    for sheet in workbook.sheetnames:
        # Normalize the workbook sheet name similarly
        normalized_sheet = re.sub(r'\s+', ' ', sheet.strip()).lower()
        if normalized_input == normalized_sheet:
            return workbook[sheet]
    log_warning(f"Sheet '{sheet_name}' not found after normalization.")
    return None

def extract_excel_value(workbook, sheet_name, col_letter, row):
    try:
        col = column_index_from_string(col_letter)
        row = int(row)
        if row < 1:
            raise ValueError("Row must be >= 1")
        sheet = find_sheet(workbook, sheet_name)
        if sheet is None:
            return None
        value = sheet.cell(row=row, column=col).value
        return str(value) if value is not None else ""
    except Exception as e:
        log_warning(f"Error accessing cell {sheet_name}:{col_letter}{row} - {e}")
        return None

def extract_excel_image(workbook, sheet_name):
    """Extract the single image from the specified Excel sheet and save it to a temporary file."""
    try:
        sheet = find_sheet(workbook, sheet_name)
        if sheet is None:
            return None
        if not sheet._images:
            log_warning(f"No images found in sheet '{sheet_name}'.")
            return None
        if len(sheet._images) > 1:
            log_warning(f"Multiple images found in sheet '{sheet_name}'. Using the first one.")
        image = sheet._images[0]  # Take the first (and assumed only) image
        # Create a temporary file to store the image
        with tempfile.NamedTemporaryFile(delete=False, suffix='.png') as tmp_file:
            tmp_file.write(image._data())
            tmp_file_path = tmp_file.name
        log_info(f"Extracted image from sheet '{sheet_name}' to temporary file '{tmp_file_path}'")
        return tmp_file_path
    except Exception as e:
        log_warning(f"Error extracting image from sheet '{sheet_name}': {e}")
        return None

def replace_placeholders_in_text(text, workbook):
    """Replace placeholders in text, handling cell references and image placeholders with flexible spacing."""
    pattern = r"\{([^}]+)\}"  # Match anything inside { }
    matches = list(re.finditer(pattern, text))
    if not matches:
        return text, None  # Return text and None for no image replacement

    new_text = ""
    last_end = 0
    image_replacement = None  # Store info for image replacement if needed

    for match in matches:
        content = match.group(1).strip()  # Strip leading/trailing spaces from content
        # Split on colon, allowing spaces around it (e.g., "Sheet1 : A1" or "Sheet1: A1")
        parts = re.split(r'\s*:\s*', content)
        if len(parts) == 2:
            sheet_name = re.sub(r'\s+', ' ', parts[0].strip())  # Normalize spaces in sheet name
            identifier = re.sub(r'\s+', '', parts[1].strip())  # Remove all spaces in identifier
            if identifier.upper() == "IMG":
                # Handle image placeholder (case-insensitive)
                image_path = extract_excel_image(workbook, sheet_name)
                if image_path:
                    image_replacement = {'sheet_name': sheet_name, 'image_path': image_path}
                    value = ""  # Replace placeholder with empty string in text
                else:
                    value = match.group(0)  # Keep placeholder if image extraction fails
            else:
                # Handle cell reference (e.g., A1, A 1)
                cell_match = re.match(r"([A-Za-z]+)(\d+)", identifier, re.IGNORECASE)
                if cell_match and cell_match.group(0) == identifier:
                    col_letter, row = cell_match.groups()
                    value = extract_excel_value(workbook, sheet_name, col_letter, row)
                    if value is None:
                        value = match.group(0)  # Keep placeholder if value couldn't be retrieved
                else:
                    log_warning(f"Invalid cell ID '{identifier}' in placeholder '{match.group(0)}'")
                    value = match.group(0)
        else:
            log_warning(f"Invalid placeholder '{match.group(0)}'")
            value = match.group(0)

        new_text += text[last_end:match.start()] + value
        last_end = match.end()

    new_text += text[last_end:]
    return new_text, image_replacement

def process_shapes(shapes, workbook, slide):
    """Process shapes in a slide, handling text boxes, tables, and image replacements."""
    shapes_to_delete = []  # Track shapes to delete (text boxes with images)
    images_to_add = []    # Track images to add with their positions

    for shape in shapes:
        if shape.has_text_frame:
            # Handle text boxes: replace placeholders or prepare image replacement
            for paragraph in shape.text_frame.paragraphs:
                original = paragraph.text
                updated, image_replacement = replace_placeholders_in_text(original, workbook)
                if original != updated:
                    log_info(f"Updated textbox from '{original}' to '{updated}'")
                    paragraph.text = updated
                if image_replacement:
                    # Store shape for deletion and image for addition
                    shapes_to_delete.append(shape)
                    images_to_add.append({
                        'image_path': image_replacement['image_path'],
                        'left': shape.left,
                        'top': shape.top,
                        'width': shape.width,
                        'height': shape.height
                    })
        elif shape.has_table:
            # Handle tables: replace placeholders within each paragraph
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        original = paragraph.text
                        updated, _ = replace_placeholders_in_text(original, workbook)  # Ignore image replacements in tables
                        if original != updated:
                            log_info(f"Updated table cell paragraph from '{original}' to '{updated}'")
                            paragraph.text = updated

    # Delete text boxes and add images after processing shapes to avoid iteration issues
    for shape in shapes_to_delete:
        # Remove the shape (text box) from the slide
        shape.text_frame.clear()  # Clear text
        shape._element.getparent().remove(shape._element)  # Remove shape from slide
        log_info(f"Removed text box shape for image replacement")

    for image_info in images_to_add:
        # Add the image to the slide at the text box's position
        try:
            img = slide.shapes.add_picture(
                image_info['image_path'],
                image_info['left'],
                image_info['top'],
                width=image_info['width'],  # Use text box width
                height=None  # Let height scale proportionally
            )
            log_info(f"Added image from '{image_info['image_path']}' at position ({image_info['left']}, {image_info['top']})")
            # Clean up temporary image file
            os.unlink(image_info['image_path'])
            log_info(f"Deleted temporary image file '{image_info['image_path']}'")
        except Exception as e:
            log_warning(f"Failed to add image from '{image_info['image_path']}': {e}")

def process_presentation(pptx_path, excel_path, output_path):
    log_info(f"Loading Excel file: {excel_path}")
    try:
        workbook = openpyxl.load_workbook(excel_path, data_only=True)
        log_info(f"Loaded Excel. Sheets: {workbook.sheetnames}")
    except Exception as e:
        log_warning(f"Failed to load Excel file: {e}")
        return

    log_info(f"Loading PowerPoint file: {pptx_path}")
    try:
        presentation = Presentation(pptx_path)
        log_info("PowerPoint file loaded")
    except Exception as e:
        log_warning(f"Failed to load PowerPoint file: {e}")
        return

    for slide_index, slide in enumerate(presentation.slides, 1):
        log_info(f"Processing slide {slide_index}")
        process_shapes(slide.shapes, workbook, slide)

    log_info(f"Saving updated presentation to: {output_path}")
    try:
        presentation.save(output_path)
        log_info("Presentation saved successfully")
    except Exception as e:
        log_warning(f"Failed to save output file: {e}")

if __name__ == "__main__":
    pptx_file = "template.pptx"
    excel_file = "SCFT MODEL REPORT-Updated.xlsx"
    output_file = "output.pptx"
    process_presentation(pptx_file, excel_file, output_file)